"""Versioned XLSX imports for systems, policies, and their owned children.

The module deliberately keeps the preview in memory.  No uploaded workbook or
import-batch model is persisted; a confirmation performs a fresh preview and
writes the complete workbook in one database transaction.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO
from ipaddress import ip_address, ip_interface, ip_network

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone
from ipam.models import IPAddress, IPRange, Prefix, VRF
from netbox.config import get_config
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill

from ..models import AccessPolicy, ActiveStatusChoices, ApplicationSystem, PolicyService, SystemAddress, SystemAlias
from .ipam_writes import OPERATION_CREATE, TYPE_IP_ADDRESS, TYPE_IP_RANGE, TYPE_PREFIX, write_system_address
from .policy_uniqueness import find_duplicate_policies


TEMPLATE_VERSION = "1.0"
SYSTEM_SHEETS = {
    "业务系统": ("system_name", "status", "description", "comments"),
    "系统别名": ("system_name", "alias"),
    "系统地址": ("system_name", "object_type", "address", "range_end", "vrf_rd", "native_status", "native_description"),
}
POLICY_SHEETS = {
    "访问关系": ("policy_name", "category", "source_system", "target_system", "valid_from", "valid_until", "description", "comments"),
    "服务项": ("policy_name", "protocol", "destination_port"),
}

SYSTEM_FIELD_GUIDE = (
    ("业务系统", "system_name", "必填；业务系统规范名。规范名与所有系统别名共用不区分大小写的唯一命名空间。", "Billing"),
    ("业务系统", "status", "可选；active（启用）或 inactive（停用），留空默认为 active。", "active"),
    ("业务系统", "description", "可选；简短说明。", "账单与支付系统"),
    ("业务系统", "comments", "可选；详细备注。", "由财务平台团队维护"),
    ("系统别名", "system_name", "必填；必须精确填写“业务系统”工作表中的规范名，不能填写别名。", "Billing"),
    ("系统别名", "alias", "必填；全局唯一的替代名称。一个业务系统可用多行配置多个别名。", "BILL"),
    ("系统地址", "system_name", "必填；必须精确填写业务系统规范名，不能填写别名。", "Billing"),
    ("系统地址", "object_type", "必填；从下拉框选择 prefix、ip_address 或 ip_range。", "prefix"),
    ("系统地址", "address", "必填；prefix 填 CIDR；ip_address 填带掩码的主机地址；ip_range 填起始 IP。", "10.20.0.0/24"),
    ("系统地址", "range_end", "仅 ip_range 必填并填写结束 IP；prefix 和 ip_address 必须留空。", "10.20.1.20"),
    ("系统地址", "vrf_rd", "可选；填写已有 VRF 的精确 RD。留空表示全局路由表。", "65000:10"),
    ("系统地址", "native_status", "可选；创建原生 IPAM 对象时使用。对象已存在时必须与其当前状态一致。", "active"),
    ("系统地址", "native_description", "可选；创建原生 IPAM 对象时使用。对象已存在时必须与其当前描述一致。", "Billing production subnet"),
)

SYSTEM_EXAMPLES = (
    ("业务系统", "新增业务系统", "Billing | active | 账单与支付系统 | 生产系统"),
    ("系统别名", "同一系统的两个别名（分别占一行）", "Billing | BILL；Billing | Payment Billing"),
    ("系统地址", "Prefix / 网段", "Billing | prefix | 10.20.0.0/24 | [留空] | 65000:10 | active | Billing subnet"),
    ("系统地址", "IP address / 单个地址", "Billing | ip_address | 10.20.0.15/32 | [留空] | 65000:10 | active | Billing VIP"),
    ("系统地址", "IP range / 地址范围", "Billing | ip_range | 10.20.1.10 | 10.20.1.20 | 65000:10 | active | Billing pool"),
)

POLICY_FIELD_GUIDE = (
    ("访问关系", "source_system / target_system", "填写已有业务系统的规范名或已确认别名；匹配忽略大小写但必须精确。", "Billing / CRM"),
    ("访问关系", "valid_from / valid_until", "可选；ISO 8601 时间，失效时间必须晚于生效时间。", "2026-01-01T00:00:00+08:00"),
    ("服务项", "policy_name", "必须精确引用本工作簿“访问关系”工作表中的关系名称。", "Billing to CRM"),
    ("服务项", "protocol", "必填；tcp 或 udp。", "tcp"),
    ("服务项", "destination_port", "必填；ANY、单端口（443）或闭区间（8000-8080）。", "443"),
)


@dataclass
class Issue:
    sheet: str
    row: int
    field: str
    value: object
    reason: str


@dataclass
class PlanRow:
    sheet: str
    row: int
    action: str
    summary: str
    data: dict = field(default_factory=dict, repr=False)


@dataclass
class ImportPlan:
    kind: str
    rows: list[PlanRow] = field(default_factory=list)
    errors: list[Issue] = field(default_factory=list)

    @property
    def valid(self):
        return not self.errors

    @property
    def counts(self):
        result = {}
        for row in self.rows:
            result[row.action] = result.get(row.action, 0) + 1
        return result


def build_template(kind):
    sheets = SYSTEM_SHEETS if kind == "systems" else POLICY_SHEETS
    workbook = Workbook()
    workbook.remove(workbook.active)
    info = workbook.create_sheet("模板说明")
    info.append(("template_version", TEMPLATE_VERSION))
    info.append(("import_type", kind))
    info.append(("说明", "不要修改工作表名称或表头；空白行会被忽略。导入只新增，不覆盖或删除已有对象。请先查看“字段说明”和“填写示例”工作表。"))
    info.append(("导入流程", "填写数据后上传并预检；只有全部错误修正后才能确认，确认时整个工作簿在一个事务中写入。"))
    info.append(("名称关联", "业务系统、系统别名和系统地址通过 system_name 关联；别名和地址工作表中的 system_name 必须填写规范名。访问关系可按规范名或已确认别名匹配已有系统。"))
    if kind == "policies":
        info.append(("服务项", "destination_port 仅接受 ANY、单端口（443）或闭区间（8000-8080）。每个访问关系至少一项。"))
    guide = workbook.create_sheet("字段说明")
    guide.append(("工作表", "字段", "填写规则", "示例"))
    for row in SYSTEM_FIELD_GUIDE if kind == "systems" else POLICY_FIELD_GUIDE:
        guide.append(row)
    examples = workbook.create_sheet("填写示例")
    examples.append(("工作表", "案例", "按表头顺序填写的示例（竖线分隔单元格）"))
    if kind == "systems":
        for row in SYSTEM_EXAMPLES:
            examples.append(row)
    else:
        examples.append(("访问关系", "HTTPS 访问", "Billing to CRM | integration | BILL | CRM | [留空] | [留空] | 账单调用 CRM | [留空]"))
        examples.append(("服务项", "单端口", "Billing to CRM | tcp | 443"))
        examples.append(("服务项", "端口范围", "Billing to CRM | tcp | 8000-8080"))
    for name, headers in sheets.items():
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2F6B9A")
        sheet.freeze_panes = "A2"
        if kind == "systems" and name == "系统地址":
            validation = DataValidation(type="list", formula1='"prefix,ip_address,ip_range"', allow_blank=False)
            validation.error = "请选择 prefix、ip_address 或 ip_range。"
            validation.errorTitle = "无效的地址类型"
            validation.prompt = "prefix=网段；ip_address=单个 IP；ip_range=起止范围"
            validation.promptTitle = "地址类型"
            validation.showErrorMessage = True
            validation.showInputMessage = True
            sheet.add_data_validation(validation)
            validation.add("B2:B1048576")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = max(16, len(str(column[0].value)) + 2)
    for sheet in (info, guide, examples):
        sheet.freeze_panes = "A2"
        sheet.column_dimensions["A"].width = 18
        sheet.column_dimensions["B"].width = 30
        sheet.column_dimensions["C"].width = 90
        sheet.column_dimensions["D"].width = 38
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2F6B9A")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def preview_workbook(content, kind, user):
    plan = ImportPlan(kind=kind)
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        plan.errors.append(Issue("工作簿", 0, "file", "", f"无法读取 XLSX 文件：{exc}"))
        return plan
    expected = SYSTEM_SHEETS if kind == "systems" else POLICY_SHEETS
    _validate_metadata(workbook, kind, plan)
    parsed = {}
    for sheet_name, headers in expected.items():
        parsed[sheet_name] = _read_sheet(workbook, sheet_name, headers, plan)
    if plan.errors:
        return plan
    if kind == "systems":
        _preview_systems(parsed, plan, user)
    else:
        _preview_policies(parsed, plan, user)
    return plan


def apply_workbook(content, kind, user):
    """Re-preview immediately before applying, then commit the entire plan."""
    with transaction.atomic():
        plan = preview_workbook(content, kind, user)
        if not plan.valid:
            raise ValidationError("工作簿在确认后发生冲突，请重新预检。")
        if kind == "systems":
            _apply_systems(plan, user)
        else:
            _apply_policies(plan)
    return plan


def _validate_metadata(workbook, kind, plan):
    if "模板说明" not in workbook.sheetnames:
        plan.errors.append(Issue("模板说明", 0, "template_version", "", "缺少模板说明工作表。"))
        return
    values = {str(row[0].value or "").strip(): row[1].value for row in workbook["模板说明"].iter_rows(min_row=1) if len(row) >= 2}
    if str(values.get("template_version", "")) != TEMPLATE_VERSION:
        plan.errors.append(Issue("模板说明", 1, "template_version", values.get("template_version"), f"仅支持模板版本 {TEMPLATE_VERSION}。"))
    if values.get("import_type") != kind:
        plan.errors.append(Issue("模板说明", 2, "import_type", values.get("import_type"), "模板类型与导入入口不匹配。"))


def _read_sheet(workbook, name, expected_headers, plan):
    if name not in workbook.sheetnames:
        plan.errors.append(Issue(name, 0, "sheet", "", "缺少必需工作表。"))
        return []
    sheet = workbook[name]
    first = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    headers = tuple(str(value or "").strip() for value in first)
    if headers != expected_headers:
        plan.errors.append(Issue(name, 1, "header", ", ".join(headers), f"表头必须为：{', '.join(expected_headers)}"))
        return []
    rows = []
    for number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value not in (None, "") for value in values):
            continue
        rows.append((number, {header: _clean(value) for header, value in zip(headers, values)}))
    return rows


def _clean(value):
    return value.strip() if isinstance(value, str) else value


def _error(plan, sheet, row, field, value, reason):
    plan.errors.append(Issue(sheet, row, field, value, reason))


def _resolve_system(name, pending, user):
    key = str(name or "").strip().casefold()
    if key in pending:
        return pending[key]
    systems = ApplicationSystem.objects.restrict(user, "view").filter(name__iexact=name)
    aliases = SystemAlias.objects.filter(name__iexact=name, system__in=ApplicationSystem.objects.restrict(user, "view"))
    candidates = list(systems) + [alias.system for alias in aliases]
    unique = {item.pk: item for item in candidates}
    return next(iter(unique.values())) if len(unique) == 1 else None


def _preview_systems(parsed, plan, user):
    if not user.has_perm("netbox_access_relations.add_applicationsystem"):
        _error(plan, "业务系统", 0, "permission", "", "没有新增业务系统权限。")
    if parsed["系统别名"] and not user.has_perm("netbox_access_relations.add_systemalias"):
        _error(plan, "系统别名", 0, "permission", "", "没有新增系统别名权限。")
    pending = {}
    for row, data in parsed["业务系统"]:
        name = str(data.get("system_name") or "").strip()
        status = str(data.get("status") or ActiveStatusChoices.ACTIVE).lower()
        if not name:
            _error(plan, "业务系统", row, "system_name", name, "不能为空。")
            continue
        key = name.casefold()
        if key in pending:
            _error(plan, "业务系统", row, "system_name", name, "工作簿中规范名重复。")
            continue
        existing = ApplicationSystem.objects.filter(name__iexact=name).first()
        alias_conflict = SystemAlias.objects.filter(name__iexact=name).exists()
        if alias_conflict:
            _error(plan, "业务系统", row, "system_name", name, "名称已被系统别名占用。")
        elif status not in ActiveStatusChoices.values:
            _error(plan, "业务系统", row, "status", status, "状态仅允许 active 或 inactive。")
        elif existing:
            if not ApplicationSystem.objects.restrict(user, "view").filter(pk=existing.pk).exists():
                _error(plan, "业务系统", row, "system_name", name, "对象不存在或无权查看。")
            else:
                pending[key] = existing
                plan.rows.append(PlanRow("业务系统", row, "skip", f"已存在：{name}", {**data, "existing_id": existing.pk}))
        else:
            marker = {"pending_name": name}
            pending[key] = marker
            plan.rows.append(PlanRow("业务系统", row, "create_system", f"新增业务系统：{name}", data))

    seen_aliases = set()
    for row, data in parsed["系统别名"]:
        name, alias = data.get("system_name"), str(data.get("alias") or "").strip()
        system = _resolve_system(name, pending, user)
        if not system:
            _error(plan, "系统别名", row, "system_name", name, "无法精确匹配有权查看的规范名。")
            continue
        key = alias.casefold()
        if not alias:
            _error(plan, "系统别名", row, "alias", alias, "不能为空。")
        elif key in seen_aliases:
            _error(plan, "系统别名", row, "alias", alias, "工作簿中别名重复。")
        elif ApplicationSystem.objects.filter(name__iexact=alias).exists():
            _error(plan, "系统别名", row, "alias", alias, "别名与规范名冲突。")
        else:
            seen_aliases.add(key)
            existing = SystemAlias.objects.filter(name__iexact=alias).first()
            system_name = system.get("pending_name") if isinstance(system, dict) else system.name
            if existing and not isinstance(system, dict) and existing.system_id == system.pk:
                plan.rows.append(PlanRow("系统别名", row, "skip", f"已关联别名：{alias}", data))
            elif existing:
                _error(plan, "系统别名", row, "alias", alias, "别名已属于其他业务系统。")
            else:
                plan.rows.append(PlanRow("系统别名", row, "create_alias", f"为 {system_name} 新增别名：{alias}", data))

    seen_addresses = set()
    for row, data in parsed["系统地址"]:
        _preview_address(row, data, pending, seen_addresses, plan, user)
    _validate_pending_native_address_conflicts(plan)


def _preview_address(row, data, pending, seen, plan, user):
    system = _resolve_system(data.get("system_name"), pending, user)
    native_type = str(data.get("object_type") or "").lower()
    if not system:
        _error(plan, "系统地址", row, "system_name", data.get("system_name"), "无法精确匹配有权查看的规范名。")
        return
    if native_type not in {TYPE_PREFIX, TYPE_IP_ADDRESS, TYPE_IP_RANGE}:
        _error(plan, "系统地址", row, "object_type", native_type, "仅允许 prefix、ip_address、ip_range。")
        return
    try:
        normalized = _normalize_address(native_type, data.get("address"), data.get("range_end"))
    except ValueError as exc:
        _error(plan, "系统地址", row, "address", data.get("address"), str(exc))
        return
    rd = str(data.get("vrf_rd") or "").strip()
    vrfs = VRF.objects.filter(rd=rd) if rd else VRF.objects.filter(pk__isnull=True)
    if rd and vrfs.count() != 1:
        _error(plan, "系统地址", row, "vrf_rd", rd, "VRF RD 必须精确匹配唯一 VRF。")
        return
    vrf = vrfs.first() if rd else None
    key = (str(data.get("system_name")).casefold(), native_type, normalized, rd)
    if key in seen:
        _error(plan, "系统地址", row, "address", data.get("address"), "工作簿中地址行重复。")
        return
    seen.add(key)
    model, lookup = _native_lookup(native_type, normalized, vrf)
    candidates = model.objects.filter(**lookup)
    if candidates.count() > 1:
        _error(plan, "系统地址", row, "address", data.get("address"), "存在多个精确候选，不能猜测选择。")
        return
    native = candidates.first()
    if native and not model.objects.restrict(user, "view").filter(pk=native.pk).exists():
        _error(plan, "系统地址", row, "address", data.get("address"), "对象不存在或无权查看。")
        return
    if native:
        supplied_status = str(data.get("native_status") or "").strip()
        supplied_description = str(data.get("native_description") or "").strip()
        if supplied_status and supplied_status != native.status:
            _error(plan, "系统地址", row, "native_status", supplied_status, "与已有 IPAM 对象状态冲突；导入不会更新原生对象。")
            return
        if supplied_description and supplied_description != (native.description or ""):
            _error(plan, "系统地址", row, "native_description", supplied_description, "与已有 IPAM 对象描述冲突；导入不会更新原生对象。")
            return
    system_id = None if isinstance(system, dict) else system.pk
    if native and system_id and SystemAddress.objects.filter(system_id=system_id, **{native_type: native}).exists():
        plan.rows.append(PlanRow("系统地址", row, "skip", f"已关联：{normalized}", data))
        return
    if not user.has_perm("netbox_access_relations.add_systemaddress"):
        _error(plan, "系统地址", row, "permission", "", "没有新增系统地址关联权限。")
        return
    action = "link_address" if native else "create_address"
    if not native and not user.has_perm(f"ipam.add_{model._meta.model_name}"):
        _error(plan, "系统地址", row, "permission", native_type, "没有创建此类 IPAM 对象的权限。")
        return
    data = {**data, "normalized": normalized, "vrf_id": vrf.pk if vrf else None, "native_id": native.pk if native else None}
    label = "关联已有" if native else "创建并关联"
    plan.rows.append(PlanRow("系统地址", row, action, f"{label}：{normalized}", data))


def _normalize_address(native_type, start, end):
    if not start:
        raise ValueError("地址不能为空。")
    if native_type == TYPE_PREFIX:
        if end:
            raise ValueError("Prefix 不得填写 range_end。")
        return str(ip_network(str(start), strict=False))
    if native_type == TYPE_IP_ADDRESS:
        if end:
            raise ValueError("IP 地址不得填写 range_end。")
        return str(ip_interface(str(start)))
    if not end:
        raise ValueError("IPRange 必须填写 range_end。")
    first, last = ip_address(str(start).split("/")[0]), ip_address(str(end).split("/")[0])
    if first.version != last.version or first > last:
        raise ValueError("地址范围的起止值无效。")
    return f"{first}-{last}"


def _native_lookup(native_type, normalized, vrf):
    if native_type == TYPE_PREFIX:
        return Prefix, {"prefix": normalized, "vrf": vrf}
    if native_type == TYPE_IP_ADDRESS:
        return IPAddress, {"address": normalized, "vrf": vrf}
    start, end = normalized.split("-", 1)
    return IPRange, {"start_address": start, "end_address": end, "vrf": vrf}


def _validate_pending_native_address_conflicts(plan):
    """Find native IPAM conflicts between objects created by this workbook.

    Native ``full_clean()`` checks the database, so it cannot see another
    ``create_address`` row until the import is already being applied. Check the
    subset of rules that can conflict between pending IPAddress/IPRange objects
    here, before a user can confirm the import.
    """
    pending = [
        row for row in plan.rows
        if row.sheet == "系统地址"
        and row.action == "create_address"
        and row.data["object_type"] in {TYPE_IP_ADDRESS, TYPE_IP_RANGE}
    ]
    for position, earlier in enumerate(pending):
        for later in pending[position + 1:]:
            if earlier.data["vrf_id"] != later.data["vrf_id"]:
                continue
            earlier_type = earlier.data["object_type"]
            later_type = later.data["object_type"]
            if earlier_type == later_type == TYPE_IP_RANGE:
                if _ranges_overlap(earlier.data["normalized"], later.data["normalized"]):
                    _pending_conflict_error(plan, earlier, later, "新增 IPRange 在同一 VRF 中重叠。")
            elif earlier_type == later_type == TYPE_IP_ADDRESS:
                if (
                    _ip_hosts_match(earlier.data["normalized"], later.data["normalized"])
                    and _ip_space_is_unique(earlier.data["vrf_id"])
                ):
                    _pending_conflict_error(plan, earlier, later, "新增 IPAddress 在同一 VRF 中重复。")


def _ranges_overlap(first, second):
    first_start, first_end = (ip_address(value) for value in first.split("-", 1))
    second_start, second_end = (ip_address(value) for value in second.split("-", 1))
    return first_start.version == second_start.version and first_start <= second_end and second_start <= first_end


def _ip_hosts_match(first, second):
    return ip_interface(first).ip == ip_interface(second).ip


def _ip_space_is_unique(vrf_id):
    if vrf_id is None:
        return get_config().ENFORCE_GLOBAL_UNIQUE
    return VRF.objects.get(pk=vrf_id).enforce_unique


def _pending_conflict_error(plan, first, second, reason):
    for row, other in ((first, second), (second, first)):
        _error(
            plan,
            row.sheet,
            row.row,
            "address",
            row.data["address"],
            f"{reason}与工作簿第 {other.row} 行冲突。",
        )


def _preview_policies(parsed, plan, user):
    for permission, sheet in (("add_accesspolicy", "访问关系"), ("add_policyservice", "服务项")):
        if not user.has_perm(f"netbox_access_relations.{permission}"):
            _error(plan, sheet, 0, "permission", "", "缺少导入所需权限。")
    pending = {}
    for row, data in parsed["访问关系"]:
        name = str(data.get("policy_name") or "").strip()
        key = name.casefold()
        if not name:
            _error(plan, "访问关系", row, "policy_name", name, "不能为空。")
            continue
        if key in pending:
            _error(plan, "访问关系", row, "policy_name", name, "工作簿中关系名称重复。")
            continue
        if AccessPolicy.objects.filter(name__iexact=name).exists():
            _error(plan, "访问关系", row, "policy_name", name, "同名访问关系已存在，导入不会覆盖。")
            continue
        source = _resolve_system(data.get("source_system"), {}, user)
        target = _resolve_system(data.get("target_system"), {}, user)
        if not source:
            _error(plan, "访问关系", row, "source_system", data.get("source_system"), "无法精确匹配有权查看的规范名或别名。")
        if not target:
            _error(plan, "访问关系", row, "target_system", data.get("target_system"), "无法精确匹配有权查看的规范名或别名。")
        if source and target and source.pk == target.pk:
            _error(plan, "访问关系", row, "target_system", data.get("target_system"), "源系统和目标系统必须不同。")
        try:
            valid_from = _parse_datetime(data.get("valid_from"))
            valid_until = _parse_datetime(data.get("valid_until"))
            if valid_from and valid_until and valid_until <= valid_from:
                raise ValueError("失效时间必须晚于生效时间。")
        except ValueError as exc:
            _error(plan, "访问关系", row, "valid_until", data.get("valid_until"), str(exc))
        if source and target:
            enriched = {**data, "source_id": source.pk, "target_id": target.pk, "_row": row}
            pending[key] = enriched
            plan.rows.append(PlanRow("访问关系", row, "create_policy", f"新增访问关系：{name}", enriched))

    services = {}
    for row, data in parsed["服务项"]:
        name = str(data.get("policy_name") or "").strip()
        key = name.casefold()
        if key not in pending:
            _error(plan, "服务项", row, "policy_name", name, "必须引用本工作簿中新增的访问关系。")
            continue
        protocol = str(data.get("protocol") or "").lower()
        if protocol not in {"tcp", "udp"}:
            _error(plan, "服务项", row, "protocol", protocol, "仅允许 tcp 或 udp。")
            continue
        try:
            is_any, start, end = _parse_port(data.get("destination_port"))
        except ValueError as exc:
            _error(plan, "服务项", row, "destination_port", data.get("destination_port"), str(exc))
            continue
        service_key = (protocol, is_any, start, end)
        if service_key in services.setdefault(key, set()):
            _error(plan, "服务项", row, "destination_port", data.get("destination_port"), "同一访问关系内服务项重复。")
            continue
        services[key].add(service_key)
        enriched = {**data, "protocol": protocol, "is_any": is_any, "port_start": start, "port_end": end}
        plan.rows.append(PlanRow("服务项", row, "create_service", f"{name}：{protocol.upper()}/{data.get('destination_port')}", enriched))
    for key, data in pending.items():
        if not services.get(key):
            _error(plan, "服务项", 0, "policy_name", data["policy_name"], "每个访问关系至少需要一个服务项。")
    _validate_policy_uniqueness(pending, services, plan, user)


def _validate_policy_uniqueness(pending, services, plan, user):
    """Reject policies whose endpoint + service set duplicates an existing or
    in-workbook policy.  Validity period and category are intentionally ignored."""
    seen = {}
    for key, data in pending.items():
        service_keys = frozenset(services.get(key, set()))
        if not service_keys:
            continue
        # Against policies already in the database.
        duplicates = find_duplicate_policies(
            data["source_id"], data["target_id"], service_keys, user=user,
        )
        if duplicates:
            _error(
                plan, "访问关系", data["_row"], "policy_name", data["policy_name"],
                f"已存在源系统、目标系统与服务项完全相同的访问关系：{duplicates[0].name}。",
            )
            continue
        # Against other policies created by this same workbook.
        bucket = (data["source_id"], data["target_id"], service_keys)
        if bucket in seen:
            _error(
                plan, "访问关系", data["_row"], "policy_name", data["policy_name"],
                f"与工作簿第 {seen[bucket]} 行的访问关系源系统、目标系统和服务项完全相同。",
            )
        else:
            seen[bucket] = data["_row"]


def _parse_datetime(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        result = value
    else:
        try:
            result = datetime.fromisoformat(str(value))
        except ValueError as exc:
            raise ValueError("时间必须为 ISO 8601 格式。") from exc
    return timezone.make_aware(result, timezone.get_current_timezone()) if timezone.is_naive(result) else result


def _parse_port(value):
    text = str(value or "").strip().upper()
    if text == "ANY":
        return True, None, None
    parts = text.split("-", 1)
    try:
        start = int(parts[0])
        end = int(parts[-1])
    except (TypeError, ValueError) as exc:
        raise ValueError("端口仅接受 ANY、单端口或起止范围。") from exc
    if not 1 <= start <= end <= 65535:
        raise ValueError("端口必须在 1 到 65535 之间，且起点不大于终点。")
    return False, start, end


def _apply_systems(plan, user):
    systems = {obj.name.casefold(): obj for obj in ApplicationSystem.objects.all()}
    for row in plan.rows:
        data = row.data
        if row.action == "create_system":
            obj = ApplicationSystem(name=data["system_name"], status=data.get("status") or "active", description=data.get("description") or "", comments=data.get("comments") or "")
            obj.full_clean(); obj.save()
            systems[obj.name.casefold()] = obj
        elif row.action == "create_alias":
            obj = SystemAlias(system=systems[str(data["system_name"]).casefold()], name=data["alias"])
            obj.full_clean(); obj.save()
        elif row.action in {"link_address", "create_address"}:
            system = systems[str(data["system_name"]).casefold()]
            native_type = data["object_type"]
            if row.action == "link_address":
                model = {TYPE_PREFIX: Prefix, TYPE_IP_ADDRESS: IPAddress, TYPE_IP_RANGE: IPRange}[native_type]
                native = model.objects.get(pk=data["native_id"])
                address = SystemAddress(system=system, **{native_type: native})
                address.full_clean(); address.save()
            else:
                normalized = data["normalized"]
                values = {"native_vrf": VRF.objects.filter(pk=data["vrf_id"]).first(), "native_status": data.get("native_status") or "active", "native_description": data.get("native_description") or ""}
                if native_type == TYPE_PREFIX: values["native_prefix"] = normalized
                elif native_type == TYPE_IP_ADDRESS: values["native_ip_address"] = normalized
                else: values["native_range_start"], values["native_range_end"] = normalized.split("-", 1)
                write_system_address(SystemAddress(system=system), operation=OPERATION_CREATE, native_type=native_type, native_values=values, user=user)


def _apply_policies(plan):
    policies = {}
    for row in plan.rows:
        data = row.data
        if row.action == "create_policy":
            obj = AccessPolicy(name=data["policy_name"], category=data.get("category") or "", source_system_id=data["source_id"], target_system_id=data["target_id"], valid_from=_parse_datetime(data.get("valid_from")), valid_until=_parse_datetime(data.get("valid_until")), description=data.get("description") or "", comments=data.get("comments") or "")
            obj.full_clean(); obj.save()
            policies[obj.name.casefold()] = obj
        elif row.action == "create_service":
            obj = PolicyService(policy=policies[str(data["policy_name"]).casefold()], protocol=data["protocol"], is_any=data["is_any"], port_start=data["port_start"], port_end=data["port_end"])
            obj.full_clean(); obj.save()
